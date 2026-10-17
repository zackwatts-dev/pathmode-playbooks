"""Quotes from call notes: the model reads your notes and picks items from your price list; code does the math.

    python quotes-from-notes/quotes_from_notes.py --sample --dry-run
    python quotes-from-notes/quotes_from_notes.py --notes call.md --prices price-list.csv --tax-rate 0

The model never sets a price. It chooses line items by code from your price list and a quantity, and quotes
the line of your notes each one came from. Code looks up the prices, does the arithmetic, and flags any
item that isn't on your list or isn't in your notes. You get a draft quote and a list of things to confirm
before it goes to the customer.

Writes out/quote.md, out/quote.docx (if python-docx is installed) and out/check.md. Nothing is sent.
"""

from __future__ import annotations

import json
import re
import sys
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from common.evidence import quote_found  # noqa: E402
from common.util import BadAnswer, done, model_for, parse_json, parser, read_csv, require, write_text  # noqa: E402

HERE = Path(__file__).parent
SAMPLE = HERE / "sample-data"

EXTRACT = """You turn a salesperson's call notes into the line items of a quote, using ONLY the price list
below. You never set or change a price. Answer with JSON only:
{"customer": "", "contact": "", "site": "", "scope_summary": "two or three plain sentences for the customer",
 "lines": [{"code": "a code from the price list", "quantity": number,
            "quote": "the words in the notes this comes from"}],
 "not_priced": [{"item": "something the notes need that isn't on the price list", "quote": "..."}],
 "assumptions": ["anything you assumed, e.g. a quantity the notes don't state"],
 "questions": ["what to confirm with the customer or on site"]}
Rules: a line goes in only if the notes call for it. If the notes are unsure ("probably", "check on site"),
include it, and add an assumption and a question. Things the notes say are excluded go nowhere.

Price list (code | description | unit):
{prices}"""


@dataclass
class Line:
    code: str
    description: str
    unit: str
    quantity: float
    unit_price: float

    @property
    def total(self) -> float:
        return round(self.quantity * self.unit_price, 2)


@dataclass
class Quote:
    customer: str
    contact: str
    site: str
    scope: str
    lines: list[Line] = field(default_factory=list)
    not_priced: list[str] = field(default_factory=list)
    assumptions: list[str] = field(default_factory=list)
    questions: list[str] = field(default_factory=list)
    flags: list[str] = field(default_factory=list)
    tax_rate: float = 0.0

    @property
    def subtotal(self) -> float:
        return round(sum(line.total for line in self.lines), 2)

    @property
    def tax(self) -> float:
        return round(self.subtotal * self.tax_rate, 2)

    @property
    def total(self) -> float:
        return round(self.subtotal + self.tax, 2)


def money(x: float) -> str:
    return f"${x:,.2f}"


def load_prices(path: Path) -> dict[str, dict]:
    return {row["Code"].upper(): row for row in read_csv(path)}


def build(notes: str, prices: dict[str, dict], model, tax_rate: float) -> Quote:
    table = "\n".join(f"{c} | {r['Description']} | {r['Unit']}" for c, r in prices.items())
    answer = parse_json(model.complete(EXTRACT.replace("{prices}", table), f"<<<NOTES\n{notes}\nNOTES>>>"))
    require(answer, ["lines"])
    q = Quote(
        str(answer.get("customer", "")),
        str(answer.get("contact", "")),
        str(answer.get("site", "")),
        str(answer.get("scope_summary", "")),
        tax_rate=tax_rate,
    )
    q.assumptions = [str(a) for a in answer.get("assumptions") or []]
    q.questions = [str(a) for a in answer.get("questions") or []]
    for item in answer.get("not_priced") or []:
        q.not_priced.append(str(item.get("item", item)) if isinstance(item, dict) else str(item))
    for item in answer["lines"]:
        code = str(item.get("code", "")).upper()
        row = prices.get(code)
        if row is None:
            q.flags.append(f"{code or 'an item'} isn't on your price list, so it was left off: price it yourself")
            continue
        try:
            quantity = float(item.get("quantity"))
        except (TypeError, ValueError):
            q.flags.append(f"{code}: no usable quantity ({item.get('quantity')!r}), so it was left off")
            continue
        if quantity <= 0:
            q.flags.append(f"{code}: quantity {quantity:g} isn't positive, so it was left off")
            continue
        if not quote_found(str(item.get("quote", "")), notes):
            q.flags.append(f"{code} ({row['Description']}) isn't clearly in your notes: confirm it's wanted")
        q.lines.append(Line(code, row["Description"], row["Unit"], quantity, float(row["Unit price"])))
    if not q.lines:
        q.flags.append("no line items: the notes may be too thin to quote from")
    return q


def quote_md(q: Quote, today: date) -> str:
    lines = [
        f"# Quote: {q.customer}",
        "",
        f"**Date:** {today:%B %-d, %Y}  ",
        f"**Prepared for:** {q.contact}  ",
        f"**Site:** {q.site}",
        "",
        "## Scope",
        "",
        q.scope,
        "",
        "| Item | Qty | Unit | Unit price | Amount |",
        "|---|---:|---|---:|---:|",
    ]
    lines += [
        f"| {li.description} | {li.quantity:g} | {li.unit} | {money(li.unit_price)} | {money(li.total)} |"
        for li in q.lines
    ]
    lines += ["", f"**Subtotal:** {money(q.subtotal)}  "]
    if q.tax_rate:
        lines.append(f"**Tax ({q.tax_rate:.2%}):** {money(q.tax)}  ")
    lines += [f"**Total:** {money(q.total)}", ""]
    if q.not_priced:
        lines += ["## Not included in this quote", ""] + [f"- {x}" for x in q.not_priced] + [""]
    if q.assumptions:
        lines += ["## Assumptions", ""] + [f"- {x}" for x in q.assumptions] + [""]
    lines += ["*Draft. Valid for 30 days once confirmed.*", ""]
    return "\n".join(lines)


def check_md(q: Quote) -> str:
    lines = ["# Before you send this quote", "", "Read the draft quote, then confirm each of these.", ""]
    if q.flags:
        lines += ["## Flags", ""] + [f"- [ ] {x}" for x in q.flags] + [""]
    lines += ["## Questions to confirm", ""] + [f"- [ ] {x}" for x in q.questions or ["(none listed)"]] + [""]
    lines += [
        "## Always",
        "",
        "- [ ] The price list is current.",
        "- [ ] Quantities match what you saw on site.",
        "- [ ] Tax is right for this job and location.",
        "",
    ]
    return "\n".join(lines)


def write_docx(q: Quote, path: Path, today: date) -> Path | None:
    try:
        import docx  # type: ignore
    except ImportError:
        return None
    doc = docx.Document()
    doc.add_heading(f"Quote: {q.customer}", level=1)
    doc.add_paragraph(f"Date: {today:%B %-d, %Y}\nPrepared for: {q.contact}\nSite: {q.site}")
    doc.add_heading("Scope", level=2)
    doc.add_paragraph(q.scope)
    table = doc.add_table(rows=1, cols=5)
    for cell, text in zip(table.rows[0].cells, ["Item", "Qty", "Unit", "Unit price", "Amount"], strict=True):
        cell.text = text
    for li in q.lines:
        row = table.add_row().cells
        for cell, text in zip(
            row, [li.description, f"{li.quantity:g}", li.unit, money(li.unit_price), money(li.total)], strict=True
        ):
            cell.text = text
    doc.add_paragraph(f"Subtotal: {money(q.subtotal)}")
    if q.tax_rate:
        doc.add_paragraph(f"Tax ({q.tax_rate:.2%}): {money(q.tax)}")
    doc.add_paragraph(f"Total: {money(q.total)}")
    if q.assumptions:
        doc.add_heading("Assumptions", level=2)
        for a in q.assumptions:
            doc.add_paragraph(a, style="List Bullet")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


# Offline stand-in: matches note lines to price-list words. Crude; a model reads notes far better.
OFFLINE_MATCH = [
    ("RTU-10T", r"(\w+) rooftop units?, both 10 ton|10.ton", "units"),
    ("CURB-ADPT", r"adapters?", "units"),
    ("THERM-PROG", r"thermostats?", "units"),
    ("CRANE-HALF", r"crane", 1),
    ("LAB-INST", r"(\d+) hours", "hours"),
    ("ELEC-DISC", r"disconnects?", "units"),
    ("DUCT-INSP", r"inspect", 1),
    ("DISP-OLD", r"replace both", "units"),
]


def offline(system: str, user: str) -> str:
    notes = user.split("<<<NOTES", 1)[-1]
    units = 2 if re.search(r"\b(two|2)\b.*units?|both", notes, re.I) else 1
    lines = []
    for code, pattern, qty in OFFLINE_MATCH:
        for line in notes.splitlines():
            m = re.search(pattern, line, re.I)
            if m:
                quantity = units if qty == "units" else int(m.group(1)) if qty == "hours" else qty
                lines.append({"code": code, "quantity": quantity, "quote": line.strip("- ").strip()})
                break
    who = re.search(r"Call with ([^,]+), ([^,.]+)", notes)
    site = re.search(r"at (\d+ [^,]+, \w+)", notes)
    return json.dumps(
        {
            "customer": who.group(2) if who else "",
            "contact": who.group(1) if who else "",
            "site": site.group(1) if site else "",
            "scope_summary": "Replace the rooftop units described in our call, with the related items listed below.",
            "lines": lines,
            "not_priced": [],
            "assumptions": ["Quantities as stated in the call notes."],
            "questions": ["Confirm on site whether the existing curbs fit."],
        }
    )


def main() -> None:
    p = parser("Turn call notes into a draft quote from your price list. Never sends.", SAMPLE)
    p.add_argument("--notes", type=Path)
    p.add_argument("--prices", type=Path)
    p.add_argument("--tax-rate", type=float, default=0.0, help="e.g. 0.056 for 5.6%% (default: no tax line)")
    args = p.parse_args()
    if args.sample:
        notes_path, prices_path = SAMPLE / "call-notes.md", SAMPLE / "price-list.csv"
    elif args.notes and args.prices:
        notes_path, prices_path = args.notes, args.prices
    else:
        p.error("choose --sample, or --notes and --prices")
    model = model_for(args, offline)
    try:
        q = build(notes_path.read_text(encoding="utf-8"), load_prices(prices_path), model, args.tax_rate)
    except BadAnswer as e:
        sys.exit(f"The model's answer couldn't be read ({e}). Try again, or check the notes.")
    today = date.today()
    written = [write_text(args.out / "quote.md", quote_md(q, today)), write_text(args.out / "check.md", check_md(q))]
    docx_path = write_docx(q, args.out / "quote.docx", today)
    if docx_path:
        written.append(docx_path)
    done(model, written, f"Draft quote: {len(q.lines)} lines, total {money(q.total)}. {len(q.flags)} flags to check.")


if __name__ == "__main__":
    main()
