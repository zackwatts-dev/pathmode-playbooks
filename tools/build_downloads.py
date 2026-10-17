"""Build each playbook's downloads for pathmode.com from this repository and the site's playbook pages.

    .venv/bin/python tools/build_downloads.py --site ../pathmode

For every playbook in the site's src/_data/playbooks.json, writes into the site:
- src/_includes/playbooks/<slug>-instructions.md: the instructions the page shows (from <slug>/instructions.md);
- src/assets/downloads/playbooks/<slug>/<slug>-instructions.md: the same, to download;
- .../<slug>-cost-worksheet.xlsx: the page's cost calculator as a spreadsheet with live formulas;
- .../<slug>-for-your-ai.md: the page, the instructions and the code's README in one Markdown file.

The printable guide (PDF) is made afterwards by the site's scripts/print-playbooks.mjs, from the built page.
"""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
FORMULA = re.compile(r"^[\w\s+\-*/().]+$")


def worksheet(item: dict, path: Path, checked: str) -> None:
    cost = item["cost"]
    wb = Workbook()
    ws = wb.active
    ws.title = "What it costs"
    bold, green = Font(bold=True), PatternFill("solid", fgColor="EEF5E0")
    ws["A1"] = f"{item['title']}: what it costs you"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Replace the example values in column B with your own. The results below update."
    ws["A2"].font = Font(italic=True)
    ws["A4"], ws["B4"] = "Your numbers", "Value"
    ws["A4"].font = ws["B4"].font = bold
    cells = {}
    row = 5
    for i in cost["inputs"]:
        ws.cell(row, 1, i["label"])
        c = ws.cell(row, 2, i["value"])
        c.fill = green
        cells[i["id"]] = f"B{row}"
        row += 1
    row += 1
    ws.cell(row, 1, "Results").font = bold
    row += 1
    for o in cost["outputs"]:
        formula = o["formula"]
        if not FORMULA.match(formula):
            raise ValueError(f"{item['slug']}: bad formula {formula!r}")
        names = set(re.findall(r"[A-Za-z_]\w*", formula))
        if names - cells.keys():
            raise ValueError(f"{item['slug']}: unknown inputs {names - cells.keys()} in {formula!r}")
        excel = re.sub(r"[A-Za-z_]\w*", lambda m: cells[m.group()], formula)
        ws.cell(row, 1, o["label"])
        c = ws.cell(row, 2, f"={excel}")
        c.number_format = '"$"#,##0' if o["format"] == "money" else '#,##0" hours"'
        c.font = bold
        row += 1
    row += 1
    for note in [
        f"From the PathMode playbook: https://pathmode.com/playbooks/{item['slug']}/ (checked {checked}).",
        "Measure the 'with the assistant' numbers yourself, in the playbook's test. Don't take anyone's word for them.",
    ]:
        ws.cell(row, 1, note).alignment = Alignment(wrap_text=True)
        row += 1
    ws.column_dimensions["A"].width = 78
    ws.column_dimensions["B"].width = 16
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def page_markdown(page: str, item: dict, instructions: str, repo: str) -> str:
    """The page source as plain Markdown: Nunjucks includes and variables replaced with their content."""
    body = re.sub(r"^---\n(?:.*?\n)?---\n", "", page, flags=re.S)
    calc = ["| Your numbers | Example |", "|---|---:|"]
    calc += [f"| {i['label']} | {i['value']} |" for i in item["cost"]["inputs"]]
    calc += ["", "| Result | Formula |", "|---|---|"]
    calc += [f"| {o['label']} | `{o['formula']}` |" for o in item["cost"]["outputs"]]
    body = body.replace('{% include "partials/cost-calculator.njk" %}', "\n".join(calc))
    body = re.sub(r'\{% include "playbooks/[\w-]+-instructions\.md" %\}', lambda _: instructions.strip(), body)
    body = body.replace("{{ playbooks.repo }}", repo)
    body = re.sub(r"</?details>|<summary>.*?</summary>", "", body)
    body = body.replace("](/", "](https://pathmode.com/")
    if "{%" in body or "{{" in body:
        raise ValueError(f"{item['slug']}: unreplaced template code in the page")
    return re.sub(r"\n{3,}", "\n\n", body).strip()


def for_your_ai(item: dict, page_md: str, instructions: str, readme: str, checked: str) -> str:
    return "\n".join(
        [
            f"# {item['title']}: a PathMode playbook, for your AI",
            "",
            "This file holds everything in one PathMode playbook: the page, the assistant instructions, and the README",
            "for the working code. Load it into ChatGPT, Claude, Gemini or Copilot and ask it to help you apply the",
            f"playbook to your own business. Page: https://pathmode.com/playbooks/{item['slug']}/. Vendor details "
            f"checked {checked}; menus change, so check them on screen.",
            "",
            f"**The problem:** {item['problem']}",
            "",
            f"**The playbook:** {item['summary']}",
            "",
            "---",
            "",
            page_md,
            "",
            "---",
            "",
            "# The instructions, to paste into an assistant",
            "",
            instructions.strip(),
            "",
            "---",
            "",
            "# The built version: README",
            "",
            readme.strip(),
            "",
            "---",
            "",
            "From PathMode (https://pathmode.com), by Zack Watts. Free to use and share under CC BY 4.0, with credit.",
            "",
        ]
    )


def main() -> None:
    p = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    p.add_argument("--site", type=Path, required=True, help="path to the pathmode site repository")
    args = p.parse_args()
    data = json.loads((args.site / "src/_data/playbooks.json").read_text())
    for item in data["items"]:
        slug = item["slug"]
        folder = ROOT / slug
        instructions = (folder / "instructions.md").read_text()
        readme = (folder / "README.md").read_text()
        page = (args.site / "src/playbooks" / f"{slug}.md").read_text()
        out = args.site / "src/assets/downloads/playbooks" / slug
        out.mkdir(parents=True, exist_ok=True)
        include = args.site / "src/_includes/playbooks" / f"{slug}-instructions.md"
        include.parent.mkdir(parents=True, exist_ok=True)
        include.write_text(instructions)
        (out / f"{slug}-instructions.md").write_text(instructions)
        worksheet(item, out / f"{slug}-cost-worksheet.xlsx", data["checked"])
        md = page_markdown(page, item, instructions, data["repo"])
        (out / f"{slug}-for-your-ai.md").write_text(for_your_ai(item, md, instructions, readme, data["checked"]))
        print(f"{slug}: instructions, worksheet, for-your-ai")


if __name__ == "__main__":
    main()
